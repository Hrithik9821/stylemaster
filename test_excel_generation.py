import os
import json
from fastapi.testclient import TestClient
from app import app

client = TestClient(app)

DATA_FILE = r"C:\Users\Operations1\.gemini\antigravity\scratch\jewelry_bulk_importer\styles_data.json"
OUTPUT_PATH = r"C:\Users\Operations1\Downloads\sjeplus_import_ready.xlsx"

def main():
    print("=== Testing Excel Generation ===")
    
    # 1. Load the pre-seeded style data
    if not os.path.exists(DATA_FILE):
        print(f"Error: Data file not found at {DATA_FILE}")
        return
        
    with open(DATA_FILE, "r", encoding="utf-8") as f:
        styles = json.load(f)
        
    print(f"Loaded {len(styles)} styles from draft.")
    
    # 2. Make mock API call to generate excel
    print("Calling /api/generate-excel API...")
    response = client.post("/api/generate-excel", json=styles)
    
    print(f"Response Status Code: {response.status_code}")
    print("Response JSON:")
    print(json.dumps(response.json(), indent=2))
    
    if response.status_code == 200:
        print("\nAPI call succeeded!")
        # Verify the file was created
        if os.path.exists(OUTPUT_PATH):
            print(f"Success: Output Excel file created at: {OUTPUT_PATH}")
            print(f"File Size: {os.path.getsize(OUTPUT_PATH)} bytes")
            
            # Inspect output sheet structure
            import openpyxl
            wb = openpyxl.load_workbook(OUTPUT_PATH)
            sheet = wb["Default Format"]
            print(f"Total rows in 'Default Format' sheet: {sheet.max_row}")
            
            # Print the values written to verify correct mapping
            print("\nFirst 4 rows details:")
            for r in range(1, min(5, sheet.max_row + 1)):
                row_vals = [cell.value for cell in sheet[r]]
                # Just print some key columns (e.g. SrNo, StyleCode, NetWt, ItemCode, Size, Pcs, Weight, StonePosition)
                # Let's find column indices by header name
                if r == 1:
                    headers = row_vals
                    print(f"Row {r} (Headers): count = {len(headers)}")
                else:
                    key_cols = ["SrNo", "StyleCode", "MItemCode", "NetWt", "ItemCode", "Size", "Pcs", "Weight", "StonePosition"]
                    details = []
                    for col_name in key_cols:
                        if col_name in headers:
                            idx = headers.index(col_name)
                            details.append(f"{col_name}: {row_vals[idx]}")
                    print(f"Row {r}: {', '.join(details)}")
            wb.close()
        else:
            print(f"Error: Output Excel file was not found at {OUTPUT_PATH}")
    else:
        print("API call failed!")

if __name__ == "__main__":
    main()
